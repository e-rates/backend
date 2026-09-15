from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from django.conf import settings
from django.db.models import Q
from django.utils import timezone

from .models import County, Payment
from .rate_reports import county_q

NAIROBI = ZoneInfo('Africa/Nairobi')
OPEN_STATUSES = ('pending', 'processing', 'failed')
AGING = [(30, '1–30 days'), (60, '31–60 days'), (90, '61–90 days'), (None, 'Over 90 days')]


@dataclass
class Column:
    label: str
    key: str
    kind: str = 'text'  # text | money | int | percent | date
    total: bool = False


@dataclass
class Section:
    heading: str
    columns: list
    rows: list
    note: str = ''


@dataclass
class Report:
    slug: str
    title: str
    subtitle: str
    summary: list = field(default_factory=list)  # [(label, value, kind)]
    sections: list = field(default_factory=list)
    county: str = 'All counties'


def _county_name(county):
    if not county:
        return 'All counties'
    found = County.objects.filter(county_q('name', county)).first()
    return found.name if found else county


def _local(dt):
    return dt.astimezone(NAIROBI) if dt else None


def _ward(parcel):
    return (parcel.ward or 'Unassigned').strip().title() if parcel else 'Unassigned'


def _title_ref(parcel):
    section = (parcel.props or {}).get('REG_SECTIO') if parcel else None
    return f"{section}/{parcel.parcel_ref}" if section else (parcel.parcel_ref if parcel else '')


def collections_report(year: int, county=None) -> Report:
    query = Payment.objects.filter(payment_year=year, parcel__isnull=False, is_deleted=False).exclude(status='refunded')
    if county:
        query = query.filter(county_q('parcel__county', county))
    bills = list(query.select_related('parcel'))
    wards = defaultdict(lambda: {'billed_count': 0, 'paid_count': 0, 'billed': Decimal(0), 'collected': Decimal(0)})
    monthly = defaultdict(lambda: {'count': 0, 'amount': Decimal(0)})
    for bill in bills:
        row = wards[(_ward(bill.parcel), bill.parcel.sub_county or '')]
        row['billed_count'] += 1
        row['billed'] += bill.amount
        if bill.status == 'completed':
            row['paid_count'] += 1
            row['collected'] += bill.amount
            month = _local(bill.updated_at).strftime('%Y-%m')
            monthly[month]['count'] += 1
            monthly[month]['amount'] += bill.amount

    ward_rows = []
    for (ward, sub_county), r in sorted(wards.items()):
        outstanding = r['billed'] - r['collected']
        ward_rows.append({
            'ward': ward, 'sub_county': sub_county, 'billed_count': r['billed_count'], 'paid_count': r['paid_count'],
            'billed': r['billed'], 'collected': r['collected'], 'outstanding': outstanding,
            'rate': (r['collected'] / r['billed'] * 100) if r['billed'] else Decimal(0),
        })
    billed = sum((r['billed'] for r in ward_rows), Decimal(0))
    collected = sum((r['collected'] for r in ward_rows), Decimal(0))
    month_rows = [
        {'month': datetime.strptime(m, '%Y-%m').strftime('%B %Y'), 'count': v['count'], 'amount': v['amount']}
        for m, v in sorted(monthly.items())
    ]
    return Report(
        county=_county_name(county),
        slug=f'collections-{year}',
        title=f'Land rates collections — {year}',
        subtitle=f'Rates billed and collected for the {year} rating year, by ward.',
        summary=[
            ('Bills issued', len(bills), 'int'),
            ('Amount billed', billed, 'money'),
            ('Amount collected', collected, 'money'),
            ('Outstanding', billed - collected, 'money'),
            ('Collection rate', (collected / billed * 100) if billed else Decimal(0), 'percent'),
        ],
        sections=[
            Section('Collections by ward', [
                Column('Ward', 'ward'), Column('Sub-county', 'sub_county'),
                Column('Bills', 'billed_count', 'int', True), Column('Paid', 'paid_count', 'int', True),
                Column('Billed (KES)', 'billed', 'money', True), Column('Collected (KES)', 'collected', 'money', True),
                Column('Outstanding (KES)', 'outstanding', 'money', True), Column('Collection rate', 'rate', 'percent'),
            ], ward_rows),
            Section('Collections by month paid', [
                Column('Month', 'month'), Column('Payments', 'count', 'int', True), Column('Amount (KES)', 'amount', 'money', True),
            ], month_rows, note='Grouped by the date the payment was confirmed.'),
        ],
    )


def arrears_report(as_of: date, county=None) -> Report:
    cutoff = datetime.combine(as_of, time.max, tzinfo=NAIROBI)
    query = Payment.objects.filter(status__in=OPEN_STATUSES, deadline__lt=cutoff, parcel__isnull=False, is_deleted=False)
    if county:
        query = query.filter(county_q('parcel__county', county))
    bills = list(query.select_related('parcel', 'user').order_by('parcel__ward', 'deadline'))
    rows = []
    buckets = defaultdict(lambda: {'count': 0, 'amount': Decimal(0)})
    ward_totals = defaultdict(lambda: {'count': 0, 'amount': Decimal(0), 'oldest': 0})
    for bill in bills:
        days = (cutoff - bill.deadline).days
        bucket = next(label for limit, label in AGING if limit is None or days <= limit)
        ward = _ward(bill.parcel)
        rows.append({
            'ward': ward, 'plot': bill.parcel.parcel_ref, 'title_ref': _title_ref(bill.parcel),
            'owner': bill.user.username, 'phone': bill.user.phone or '', 'email': bill.user.email,
            'year': bill.payment_year, 'amount': bill.amount, 'deadline': _local(bill.deadline).date(),
            'days': days, 'bucket': bucket,
        })
        buckets[bucket]['count'] += 1
        buckets[bucket]['amount'] += bill.amount
        ward_totals[ward]['count'] += 1
        ward_totals[ward]['amount'] += bill.amount
        ward_totals[ward]['oldest'] = max(ward_totals[ward]['oldest'], days)
    total = sum((r['amount'] for r in rows), Decimal(0))
    return Report(
        county=_county_name(county),
        slug=f'arrears-{as_of.isoformat()}',
        title='Defaulters and arrears',
        subtitle=f'Unpaid rate bills past their deadline as of {as_of.strftime("%d %B %Y")}.',
        summary=[
            ('Overdue bills', len(rows), 'int'),
            ('Defaulters', len({r['owner'] for r in rows}), 'int'),
            ('Total arrears', total, 'money'),
        ],
        sections=[
            Section('Arrears by age', [
                Column('Overdue for', 'bucket'), Column('Bills', 'count', 'int', True), Column('Amount (KES)', 'amount', 'money', True),
            ], [{'bucket': label, **buckets[label]} for _, label in AGING]),
            Section('Arrears by ward', [
                Column('Ward', 'ward'), Column('Bills', 'count', 'int', True),
                Column('Amount (KES)', 'amount', 'money', True), Column('Oldest (days)', 'oldest', 'int'),
            ], [{'ward': w, **v} for w, v in sorted(ward_totals.items(), key=lambda kv: -kv[1]['amount'])]),
            Section('Defaulters', [
                Column('Ward', 'ward'), Column('Plot', 'plot'), Column('Title ref.', 'title_ref'),
                Column('Owner', 'owner'), Column('Phone', 'phone'), Column('Rating year', 'year'),
                Column('Amount (KES)', 'amount', 'money', True), Column('Deadline', 'deadline', 'date'),
                Column('Days overdue', 'days', 'int'),
            ], rows),
        ],
    )


def register_report(start: date, end: date, county=None) -> Report:
    since = datetime.combine(start, time.min, tzinfo=NAIROBI)
    until = datetime.combine(end, time.max, tzinfo=NAIROBI)
    query = Payment.objects.filter(status='completed', updated_at__range=(since, until), is_deleted=False)
    if county:
        query = query.filter(county_q('parcel__county', county) | (Q(parcel__isnull=True) & county_q('user__county', county)))
    payments = list(query.select_related('parcel', 'user').order_by('updated_at'))
    rows = []
    for p in payments:
        receipt = p.processor_ref if p.processor_ref and not p.processor_ref.startswith('ws_CO_') else ''
        rows.append({
            'paid_at': _local(p.updated_at).strftime('%Y-%m-%d %H:%M'), 'receipt': receipt or 'Pending reconciliation',
            'plot': p.parcel.parcel_ref if p.parcel_id else '', 'ward': _ward(p.parcel) if p.parcel_id else '',
            'payer': p.user.username, 'phone': (p.metadata or {}).get('mpesa_payer_phone') or p.user.phone or '',
            'year': p.payment_year or '', 'channel': (p.processor or 'manual').upper() if p.processor else 'Manual',
            'amount': p.amount,
        })
    total = sum((r['amount'] for r in rows), Decimal(0))
    unreconciled = sum(1 for r in rows if r['receipt'] == 'Pending reconciliation')
    period = f'{start.strftime("%d %b %Y")} – {end.strftime("%d %b %Y")}'
    return Report(
        county=_county_name(county),
        slug=f'payment-register-{start.isoformat()}-to-{end.isoformat()}',
        title='Payment register',
        subtitle=f'Confirmed rate payments received {period}, for reconciliation against the M-Pesa statement.',
        summary=[
            ('Payments', len(rows), 'int'),
            ('Total received', total, 'money'),
            ('Without M-Pesa receipt', unreconciled, 'int'),
        ],
        sections=[Section('Payments received', [
            Column('Paid at', 'paid_at'), Column('M-Pesa receipt', 'receipt'), Column('Plot', 'plot'), Column('Ward', 'ward'),
            Column('Payer', 'payer'), Column('Paying phone', 'phone'), Column('Rating year', 'year'),
            Column('Channel', 'channel'), Column('Amount (KES)', 'amount', 'money', True),
        ], rows, note='"Pending reconciliation" means the payment was confirmed by status check and has no receipt code.')],
    )


def _fmt(value, kind):
    if value is None or value == '':
        return ''
    if kind == 'money':
        return f'{Decimal(value):,.2f}'
    if kind == 'percent':
        return f'{Decimal(value):.1f}%'
    if kind == 'int':
        return f'{int(value):,}'
    if kind == 'date':
        return value.strftime('%d %b %Y') if hasattr(value, 'strftime') else str(value)
    return str(value)


def _generated_at():
    return timezone.now().astimezone(NAIROBI).strftime('%d %B %Y, %H:%M')


def render_xlsx(report: Report, generated_by: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    head_fill = PatternFill('solid', fgColor='EDEDED')
    rule = Border(bottom=Side(style='thin', color='BDBDBD'))
    formats = {'money': '#,##0.00', 'int': '#,##0', 'percent': '0.0"%"', 'date': 'dd mmm yyyy'}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = report.county
    ws['A1'].font = Font(bold=True, size=12)
    ws['A2'] = report.title
    ws['A2'].font = Font(bold=True, size=16)
    ws['A3'] = report.subtitle
    ws['A4'] = f'Generated {_generated_at()} by {generated_by}'
    ws['A4'].font = Font(italic=True, color='737373')
    for i, (label, value, kind) in enumerate(report.summary, start=6):
        ws.cell(i, 1, label).font = bold
        cell = ws.cell(i, 2, float(value) if isinstance(value, Decimal) else value)
        cell.number_format = formats.get(kind, 'General')
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 22

    for section in report.sections:
        sheet = wb.create_sheet(section.heading[:31])
        for c, col in enumerate(section.columns, start=1):
            cell = sheet.cell(1, c, col.label)
            cell.font, cell.fill, cell.border = bold, head_fill, rule
            cell.alignment = Alignment(horizontal='right' if col.kind in formats else 'left')
        for r, row in enumerate(section.rows, start=2):
            for c, col in enumerate(section.columns, start=1):
                value = row.get(col.key)
                if isinstance(value, Decimal):
                    value = float(value)
                cell = sheet.cell(r, c, value)
                if col.kind in formats:
                    cell.number_format = formats[col.kind]
        last = len(section.rows) + 1
        if section.rows and any(col.total for col in section.columns):
            total_row = last + 1
            sheet.cell(total_row, 1, 'Total').font = bold
            for c, col in enumerate(section.columns, start=1):
                if col.total:
                    letter = get_column_letter(c)
                    cell = sheet.cell(total_row, c, f'=SUM({letter}2:{letter}{last})')
                    cell.font, cell.number_format = bold, formats[col.kind]
        if section.note:
            sheet.cell(last + 3, 1, section.note).font = Font(italic=True, color='737373')
        for c, col in enumerate(section.columns, start=1):
            longest = max([len(col.label)] + [len(_fmt(row.get(col.key), col.kind)) for row in section.rows])
            sheet.column_dimensions[get_column_letter(c)].width = min(max(longest + 2, 10), 48)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(section.columns))}{last}'

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pdf(report: Report, generated_by: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ink, muted, rule, band = colors.HexColor('#171717'), colors.HexColor('#737373'), colors.HexColor('#d4d4d4'), colors.HexColor('#f5f5f5')
    base = ParagraphStyle('base', fontName='Helvetica', fontSize=8.5, leading=11, textColor=ink)
    small = ParagraphStyle('small', parent=base, fontSize=7.5, textColor=muted)
    h1 = ParagraphStyle('h1', parent=base, fontName='Helvetica-Bold', fontSize=16, leading=20)
    h2 = ParagraphStyle('h2', parent=base, fontName='Helvetica-Bold', fontSize=11, leading=14, spaceBefore=10, spaceAfter=4)
    generated = _generated_at()

    def build(total_pages):
        def frame(canvas, doc):
            canvas.saveState()
            width, height = landscape(A4)
            canvas.setFont('Helvetica-Bold', 9)
            canvas.setFillColor(ink)
            canvas.drawString(15 * mm, height - 11 * mm, report.county.upper())
            canvas.setFont('Helvetica', 8)
            canvas.setFillColor(muted)
            canvas.drawRightString(width - 15 * mm, height - 11 * mm, report.title)
            canvas.setStrokeColor(rule)
            canvas.line(15 * mm, height - 13 * mm, width - 15 * mm, height - 13 * mm)
            canvas.drawString(15 * mm, 9 * mm, f'Generated {generated} by {generated_by} · E-Rates')
            canvas.drawRightString(width - 15 * mm, 9 * mm, f'Page {doc.page} of {total_pages}')
            canvas.restoreState()

        out = BytesIO()
        doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                                topMargin=20 * mm, bottomMargin=16 * mm, title=report.title, author=report.county)
        story = [Paragraph(report.title, h1), Paragraph(report.subtitle, small), Spacer(1, 6 * mm)]
        summary = Table(
            [[Paragraph(label, small) for label, _, _ in report.summary],
             [Paragraph(f'<b>{_fmt(v, k)}</b>', ParagraphStyle('sv', parent=base, fontSize=12, leading=15)) for _, v, k in report.summary]],
            hAlign='LEFT',
        )
        summary.setStyle(TableStyle([
            ('BOX', (0, 0), (-1, -1), 0.5, rule), ('INNERGRID', (0, 0), (-1, -1), 0.5, rule),
            ('LEFTPADDING', (0, 0), (-1, -1), 8), ('RIGHTPADDING', (0, 0), (-1, -1), 12),
            ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story += [summary, Spacer(1, 4 * mm)]
        for section in report.sections:
            numeric = [i for i, c in enumerate(section.columns) if c.kind in ('money', 'int', 'percent')]
            data = [[c.label for c in section.columns]]
            data += [[_fmt(row.get(c.key), c.kind) for c in section.columns] for row in section.rows]
            if not section.rows:
                data.append(['No records'] + [''] * (len(section.columns) - 1))
            has_total = bool(section.rows) and any(c.total for c in section.columns)
            if has_total:
                data.append([
                    'Total' if i == 0 else (_fmt(sum((row.get(c.key) or 0 for row in section.rows), Decimal(0)), c.kind) if c.total else '')
                    for i, c in enumerate(section.columns)
                ])
            avail = landscape(A4)[0] - 30 * mm - 12  # SimpleDocTemplate frame padding
            weights = [max(len(str(r[i])) for r in data) + 4 for i in range(len(section.columns))]
            table = Table(data, repeatRows=1, hAlign='LEFT', colWidths=[avail * w / sum(weights) for w in weights])
            style = [
                ('FONT', (0, 0), (-1, -1), 'Helvetica', 8), ('TEXTCOLOR', (0, 0), (-1, -1), ink),
                ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 7.5), ('TEXTCOLOR', (0, 0), (-1, 0), muted),
                ('BACKGROUND', (0, 0), (-1, 0), band), ('LINEBELOW', (0, 0), (-1, 0), 0.6, rule),
                ('LINEBELOW', (0, 1), (-1, -1), 0.25, rule),
                ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ] + [('ALIGN', (i, 0), (i, -1), 'RIGHT') for i in numeric]
            if has_total:
                style += [('FONT', (0, -1), (-1, -1), 'Helvetica-Bold', 8), ('LINEABOVE', (0, -1), (-1, -1), 0.8, ink)]
            table.setStyle(TableStyle(style))
            story += [Paragraph(section.heading, h2), table]
            if section.note:
                story += [Spacer(1, 1.5 * mm), Paragraph(section.note, small)]
        doc.build(story, onFirstPage=frame, onLaterPages=frame)
        return out.getvalue(), doc.page

    _, pages = build('?')
    pdf, _ = build(pages)
    return pdf


BUILDERS = {'collections': collections_report, 'arrears': arrears_report, 'register': register_report}
