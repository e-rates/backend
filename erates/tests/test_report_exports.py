from datetime import timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from django.contrib.gis.geos import Polygon
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from erates.models import Parcel, Payment, User
from erates.payment_flow import generate_rate_bills


class ReportExportTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user('admin', 'admin@example.com', 'Password123!', role='admin')
        owner = User.objects.create_user('owner', 'owner@example.com', 'Password123!', phone='0711111111')
        for i, ward in enumerate(['karura', 'karura', 'mugunda']):
            Parcel.objects.create(
                parcel_ref=f'P{i}', geom=Polygon.from_bbox((36.8 + i / 100, -1.3, 36.801 + i / 100, -1.299)),
                county='Nyeri', sub_county='Tetu', ward=ward, owner_user=owner, props={'REG_SECTIO': 'AGUTHI-GAAKI'},
            )
        generate_rate_bills(2026, timezone.now() + timedelta(days=30))
        paid = Payment.objects.get(parcel__parcel_ref='P0')
        paid.status, paid.processor, paid.processor_ref = 'completed', 'mpesa', 'RCPT1'
        paid.save()
        Payment.objects.filter(parcel__parcel_ref='P2').update(deadline=timezone.now() - timedelta(days=45))
        self.client.force_authenticate(self.admin)

    def fetch(self, query):
        resp = self.client.get(f'/api/reports/download/?{query}')
        self.assertEqual(resp.status_code, 200, resp.content[:300])
        return resp

    def test_collections_excel_has_ward_rows_and_formula_totals(self):
        resp = self.fetch('report=collections&file_format=xlsx&year=2026')
        self.assertIn('collections-2026.xlsx', resp['Content-Disposition'])
        wb = load_workbook(BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ['Summary', 'Collections by ward', 'Collections by month paid'])
        wards = wb['Collections by ward']
        self.assertEqual([wards.cell(r, 1).value for r in (2, 3, 4)], ['Karura', 'Mugunda', 'Total'])
        self.assertEqual(wards.cell(4, 3).value, '=SUM(C2:C3)')

    def test_arrears_buckets_overdue_bill(self):
        wb = load_workbook(BytesIO(self.fetch('report=arrears&file_format=xlsx').content))
        aging = {wb['Arrears by age'].cell(r, 1).value: wb['Arrears by age'].cell(r, 2).value for r in range(2, 6)}
        self.assertEqual(aging['31–60 days'], 1)
        defaulters = wb['Defaulters']
        self.assertEqual((defaulters.cell(2, 2).value, defaulters.cell(2, 3).value), ('P2', 'AGUTHI-GAAKI/P2'))

    def test_register_lists_receipts(self):
        today = timezone.now().astimezone(ZoneInfo('Africa/Nairobi')).date().isoformat()
        wb = load_workbook(BytesIO(self.fetch(f'report=register&file_format=xlsx&from={today}&to={today}').content))
        sheet = wb['Payments received']
        self.assertEqual(sheet.cell(2, 2).value, 'RCPT1')

    def test_every_report_renders_pdf(self):
        for report in ('collections', 'arrears', 'register'):
            resp = self.fetch(f'report={report}&file_format=pdf')
            self.assertEqual(resp['Content-Type'], 'application/pdf')
            self.assertTrue(resp.content.startswith(b'%PDF'), report)
            self.assertNotIn(b'of ?', resp.content)

    def test_validation_and_permissions(self):
        self.assertEqual(self.client.get('/api/reports/download/?report=nope&file_format=pdf').status_code, 400)
        self.assertEqual(self.client.get('/api/reports/download/?report=register&file_format=pdf&from=2026-02-01&to=2026-01-01').status_code, 400)
        self.client.force_authenticate(User.objects.get(username='owner'))
        self.assertEqual(self.client.get('/api/reports/download/?report=collections&file_format=pdf').status_code, 403)
